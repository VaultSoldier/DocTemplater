import logging
import os
import random
import sys
import tempfile
import threading
from io import BytesIO
from multiprocessing import Process, Queue
from typing import Final, Iterable, List, Optional

from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer
from docxtpl import DocxTemplate

from core.processing.data import SqliteData, get_resource_path_temp
from core.types import QuestionType


class DocxProcessingError(Exception):
    """Базовое исключение."""

    pass


class NoQuestionsError(DocxProcessingError):
    pass


class InvalidNumberError(DocxProcessingError):
    pass


class UnknownTicketTypeError(DocxProcessingError):
    pass


class DocxProcessingCancel(DocxProcessingError):
    pass


def get_list_safe(items: list, items_index: int, is_fallback: Optional[bool] = False) -> str:
    if not items:
        return ""
    try:
        return str(items[items_index])
    except IndexError:
        if is_fallback:
            return str(random.choice(items))
        return ""


def get_question(questions: List[str], status_rnd: str, index_question: int):
    if questions == []:
        return ""

    match status_rnd:
        case "fallback":
            return get_list_safe(items=questions, items_index=index_question, is_fallback=True)
        case "always":
            return str(random.choice(questions))
        case "none":
            return get_list_safe(items=questions, items_index=index_question)
        case _:
            return ""


def _batch_worker(
    batch: range,
    tpl_path: str,
    context: dict,
    output_path: str,
    status_rnd_practical: str,
    status_rnd_theoretical: str,
    practical_questions: list,
    theoretical_questions: list,
    result_queue: Queue,
):
    try:
        tpl = DocxTemplate(tpl_path)
        buffers = []

        for i in batch:
            buf = BytesIO()

            context_data = {
                "ticket_num": f"{i + 1}",
                "question_theoretical": get_question(
                    theoretical_questions, status_rnd_theoretical, i
                ),
                "question_practical": get_question(
                    practical_questions, status_rnd_practical, i
                ),
            }
            context.update(context_data)
            tpl.render(context)
            tpl.save(buf)
            buf.seek(0)
            buffers.append(buf)

        master = Document(buffers[0])
        composer = Composer(master)
        master.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        for idx, buf in enumerate(buffers[1:], start=1):
            doc = Document(buf)
            composer.append(doc)
            if idx != len(buffers) - 1:
                master.add_page_break()
        composer.save(output_path)

        result_queue.put({"ok": True})
    except Exception as e:
        result_queue.put({"ok": False, "error": str(e)})


class Processing:
    def __init__(self) -> None:
        self.PATH_BASE_DOC: Final[str] = get_resource_path_temp(
            "assets/templates/base.docx"
        )
        self.sql = SqliteData()
        self.is_windows = sys.platform.startswith("win")

        self.questions_practical: list[str] = []
        self.practical_questions_count: int = 0

        self.questions_theoretical: list[str] = []
        self.theoretical_questions_count: int = 0

        self._cancel_event = threading.Event()

    def cancel(self):
        self._cancel_event.set()

    def _check_cancel(self):
        if self._cancel_event.is_set():
            raise DocxProcessingCancel("Отмена генерации")

    def questions_import(self):
        self.questions_practical: list[str] = self.sql.read_questions_list(
            QuestionType.PRACTICAL
        )
        self.practical_questions_count: int = len(self.questions_practical)

        self.questions_theoretical: list[str] = self.sql.read_questions_list(
            QuestionType.THEORETICAL
        )
        self.theoretical_questions_count: int = len(self.questions_theoretical)

    def process_docx(
        self,
        save_to: str,
        subject: Optional[str],
        spec: Optional[str],
        cmk: Optional[str],
        tutor: Optional[str],
        date: list,
        tickets_count: int | None,
        qualify_status: bool | None,
        tickets_count_type: str,
        theoretical_rnd_type: str,
        practical_rnd_type: str,
        batch_size: int = 50,  # 50/999 * 2.5GB ≈ 125MB
    ) -> None:
        self._cancel_event.clear()  # reset cancel status before run

        logging.info(
            f"subject: {subject}\nspec: {spec}\ncmk: {cmk}\ntutor: {tutor}\ndate: {date}\n"
        )

        #  INFO: ОБНОВЛЕНИЕ ВОПРОСОВ
        self.questions_import()

        match tickets_count_type:
            case "Manual" if tickets_count is None or tickets_count <= 0:
                raise InvalidNumberError("Неверное количество билетов")
            case "Practical" if self.practical_questions_count <= 0:
                raise NoQuestionsError("Нету практических вопросов")
            case "Theoretical" if self.theoretical_questions_count <= 0:
                raise NoQuestionsError("Нету теоретических вопросов")
            case "Manual" if tickets_count is not None:
                tickets = range(tickets_count)
            case "Practical":
                tickets = range(self.practical_questions_count)
            case "Theoretical":
                tickets = range(self.theoretical_questions_count)
            case _:
                raise UnknownTicketTypeError(
                    f"Неизвестный тип билетов: {tickets_count_type}"
                )

        TOTAL_CMK_WIDTH: Final[int] = 23

        if cmk and cmk.strip():
            cmk_padding = "_" * (TOTAL_CMK_WIDTH - len(str(cmk)))
            cmk = f"{cmk_padding}{cmk}"
        else:
            cmk = "_" * (TOTAL_CMK_WIDTH - 9)

        qualify = " (квалификационный)" if qualify_status else ""
        day = f"{int(date[2]):02}" or "__"  # add "0" to single num (1 = 01, 10 = 10)
        month = date[1] or ""
        year = date[0] or ""

        context = {
            "qualify": qualify,
            "subject": subject,
            "spec": spec,
            "cmk": cmk,
            "tutor": tutor,
            "day": day,
            "month": month,
            "year": year,
        }
        tpl = DocxTemplate(self.PATH_BASE_DOC)

        if len(tickets) == 1:
            num_of_tickets: int = tickets[0]
            question_theoretical = get_question(
                questions=self.questions_theoretical,
                status_rnd=theoretical_rnd_type,
                index_question=num_of_tickets,
            )
            question_practical = get_question(
                questions=self.questions_practical,
                status_rnd=practical_rnd_type,
                index_question=num_of_tickets,
            )
            context_extend = {
                "ticket_num": str(num_of_tickets + 1),
                "question_theoretical": question_theoretical,
                "question_practical": question_practical,
            }
            context.update(context_extend)
            tpl.render(context)
            tpl.save(save_to)
            return

        context_extend = {
            "ticket_num": "{{ticket_num}}",
            "question_theoretical": "{{question_theoretical}}",
            "question_practical": "{{question_practical}}",
        }
        context.update(context_extend)

        tmp_base_docx_file = tempfile.NamedTemporaryFile(
            prefix="tmp_base_",
            suffix=".docx",
            delete=not self.is_windows,
        )

        tpl.render(context)
        tpl.save(tmp_base_docx_file.name)

        batch_files = []
        try:
            for batch_start in range(0, len(tickets), batch_size):
                self._check_cancel()
                batch = range(batch_start, min(batch_start + batch_size, len(tickets)))

                batch_file = tempfile.NamedTemporaryFile(
                    prefix=f"batch_{batch_start}_",
                    suffix=".docx",
                    delete=not self.is_windows,
                )

                result_queue = Queue()
                process = Process(
                    target=_batch_worker,
                    kwargs={
                        "batch": batch,
                        "tpl_path": tmp_base_docx_file.name,
                        "context": context.copy(),
                        "output_path": batch_file.name,
                        "status_rnd_practical": practical_rnd_type,
                        "status_rnd_theoretical": theoretical_rnd_type,
                        "practical_questions": self.questions_practical,
                        "theoretical_questions": self.questions_theoretical,
                        "result_queue": result_queue,
                    },
                )
                process.start()
                process.join()  # subprocess fully exits here, OS reclaims all lxml memory

                result = result_queue.get()
                if not result["ok"]:
                    raise DocxProcessingError(result["error"])

                batch_files.append(batch_file)

            self.merge_files(batch_files, save_to)

        except DocxProcessingError:
            threading.Thread(
                target=self.clean,
                kwargs={"path": tmp_base_docx_file, "paths": batch_files},
            ).start()
            raise

    def merge_files(
        self, files: list[tempfile._TemporaryFileWrapper], save_to: str
    ) -> None:
        """Merge temp files (batches) into final output."""
        master = Document(files[0].name)
        composer = Composer(master)
        master.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        for idx, f in enumerate(files[1:], start=1):
            doc = Document(f.name)
            composer.append(doc)
            if idx != len(files) - 1:
                master.add_page_break()

        composer.save(save_to)

    def replace_questions(
        self,
        status_rnd_practical: str,
        status_rnd_theoretical: str,
        tickets: range,
        tpl_template_file: tempfile._TemporaryFileWrapper,
    ) -> list[BytesIO]:
        """Replace questions"""

        tpl = DocxTemplate(tpl_template_file.name)
        buffers_docx = []
        try:
            for i in tickets:
                self._check_cancel()
                buffer_docx = BytesIO()

                question_theoretical = get_question(
                    questions=self.questions_theoretical,
                    status_rnd=status_rnd_theoretical,
                    index_question=i,
                )
                question_practical = get_question(
                    questions=self.questions_practical,
                    status_rnd=status_rnd_practical,
                    index_question=i,
                )
                context = {
                    "ticket_num": f"{i + 1}",
                    "question_theoretical": question_theoretical,
                    "question_practical": question_practical,
                }

                tpl.render(context)
                tpl.save(buffer_docx)
                buffer_docx.seek(0)
                buffers_docx.append(buffer_docx)
        except DocxProcessingError:
            for buf in buffers_docx:
                buf.close()
            buffers_docx.clear()

            threading.Thread(
                target=self.clean, kwargs={"path": tpl_template_file}
            ).start()
        return buffers_docx

    def docx_merge(self, buffers: list[BytesIO], save_to: str) -> None:
        """Merge temporary files"""
        master = Document(buffers[0])
        composer = Composer(master)

        # master file page break
        master.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        for idx, buf in enumerate(buffers[1:], start=1):
            doc = Document(buf)
            composer.append(doc)
            if idx != len(buffers) - 1:
                master.add_page_break()
        composer.save(save_to)

    def clean(
        self,
        path: Optional[tempfile._TemporaryFileWrapper] = None,
        paths: Optional[Iterable[tempfile._TemporaryFileWrapper]] = None,
    ):
        if path is None and paths is None:
            return

        all_paths = []
        if path is not None:
            all_paths.append(path)
        if paths is not None:
            all_paths.extend(paths)

        # clean temporary files
        for file in all_paths:
            try:
                file_path = file.name
                # Deletes files on UNIX
                # And makes sure it closes before deleting on non UNIX
                if not file.closed:
                    file.close()
                # Deletes files on non UNIX systems
                if os.path.exists(file_path):
                    os.remove(file_path)
            except Exception as e:
                logging.error(f"Failed to delete {file.name}: {e}")
